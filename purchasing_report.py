from __future__ import annotations

import argparse
import csv
from pathlib import Path
import os
import tempfile
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


REPORT_COLUMNS = (
    "status",
    "supplier",
    "sku",
    "currency",
    "previous_cost",
    "current_cost",
    "absolute_change",
    "percent_change",
    "gross_margin_percent",
    "risk",
)


class PurchasingReportError(ValueError):
    pass


def _read_rows(path: Path) -> list[dict[str, str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if tuple(reader.fieldnames or ()) != REPORT_COLUMNS:
                raise PurchasingReportError(
                    "input must be a Supplier Price Watch CSV report with the canonical columns"
                )
            return [dict(row) for row in reader]
    except OSError as exc:
        raise PurchasingReportError(f"cannot read report CSV: {path}") from exc


def _summary_rows(rows: Iterable[dict[str, str]]) -> list[tuple[str, int]]:
    materialized = list(rows)
    matched = [row for row in materialized if row["status"] == "matched"]
    return [
        ("matched", len(matched)),
        ("critical", sum(row["risk"] == "critical" for row in matched)),
        ("warning", sum(row["risk"] == "warning" for row in matched)),
        (
            "price_up",
            sum(
                (row["absolute_change"] or "0").startswith("-") is False
                and row["absolute_change"] not in ("", "0", "0.0", "0.00")
                for row in matched
            ),
        ),
        ("added", sum(row["status"] == "added" for row in materialized)),
        ("removed", sum(row["status"] == "removed" for row in materialized)),
    ]


def _autofit(sheet) -> None:
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 40)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def write_purchasing_workbook(source_csv: Path, output_xlsx: Path) -> None:
    rows = _read_rows(source_csv)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["metric", "count"])
    for row in _summary_rows(rows):
        summary.append(list(row))
    summary.freeze_panes = "A2"
    for cell in summary[1]:
        cell.font = Font(bold=True)
    _autofit(summary)

    lines = workbook.create_sheet("Purchasing Review")
    lines.append(list(REPORT_COLUMNS))
    for row in sorted(
        rows,
        key=lambda item: (
            {"critical": 0, "warning": 1, "ok": 2, "": 3}.get(item["risk"], 4),
            item["supplier"].casefold(),
            item["sku"].casefold(),
        ),
    ):
        lines.append([row[column] for column in REPORT_COLUMNS])
    lines.freeze_panes = "A2"
    lines.auto_filter.ref = lines.dimensions
    for cell in lines[1]:
        cell.font = Font(bold=True)
    _autofit(lines)

    temp_path: Path | None = None
    try:
        output_xlsx.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            dir=output_xlsx.parent,
            prefix=f".{output_xlsx.name}.",
            suffix=".xlsx",
            delete=False,
        ) as handle:
            temp_path = Path(handle.name)
        workbook.save(temp_path)
        os.replace(temp_path, output_xlsx)
        temp_path = None
    except OSError as exc:
        raise PurchasingReportError(f"cannot write XLSX report: {output_xlsx}") from exc
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert a Supplier Price Watch CSV report into an Excel purchasing workbook."
    )
    parser.add_argument("source_csv", type=Path)
    parser.add_argument("output_xlsx", type=Path)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        write_purchasing_workbook(args.source_csv, args.output_xlsx)
    except PurchasingReportError as exc:
        print(f"error: {exc}")
        return 2
    print(f"Workbook written: {args.output_xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
