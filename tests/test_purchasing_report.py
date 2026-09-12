from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from purchasing_report import PurchasingReportError, write_purchasing_workbook


class PurchasingReportTests(unittest.TestCase):
    def test_creates_summary_and_prioritized_review_sheet(self) -> None:
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "report.csv"
            output = root / "purchasing.xlsx"
            source.write_text(
                "status,supplier,sku,currency,previous_cost,current_cost,absolute_change,percent_change,gross_margin_percent,risk\n"
                "matched,Beta,B-1,TRY,100,120,20,20,18,warning\n"
                "matched,Alpha,A-1,TRY,100,140,40,40,10,critical\n"
                "added,Alpha,A-2,TRY,,90,,,,\n"
                "removed,Gamma,G-1,TRY,80,,,,,\n",
                encoding="utf-8",
            )

            write_purchasing_workbook(source, output)
            workbook = load_workbook(output, data_only=True)

            self.assertEqual(workbook.sheetnames, ["Summary", "Purchasing Review"])
            summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(summary["matched"], 2)
            self.assertEqual(summary["critical"], 1)
            self.assertEqual(summary["warning"], 1)
            self.assertEqual(summary["added"], 1)
            self.assertEqual(summary["removed"], 1)

            review = workbook["Purchasing Review"]
            self.assertEqual(review.freeze_panes, "A2")
            self.assertEqual(review["J2"].value, "critical")
            self.assertEqual(review["C2"].value, "A-1")
            self.assertTrue(review.auto_filter.ref)

    def test_rejects_noncanonical_csv(self) -> None:
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "bad.csv"
            source.write_text("sku,cost\nA,1\n", encoding="utf-8")
            with self.assertRaises(PurchasingReportError):
                write_purchasing_workbook(source, Path(temp_dir) / "out.xlsx")


if __name__ == "__main__":
    unittest.main()
