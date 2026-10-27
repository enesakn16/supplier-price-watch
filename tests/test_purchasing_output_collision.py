"""Regression coverage for source CSV preservation during XLSX export."""

import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from purchasing_report import PurchasingReportError, write_purchasing_workbook


class WorkbookOutputCollisionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "report.csv"
        self.source.write_text(
            "status,supplier,sku,currency,previous_cost,current_cost,absolute_change,percent_change,gross_margin_percent,risk\n"
            "matched,Demo,SKU-1,TRY,100,110,10,10,20,warning\n",
            encoding="utf-8",
        )

    def assert_safe_rejection(self, output: Path) -> None:
        before = self.source.read_bytes()
        with self.assertRaisesRegex(PurchasingReportError, "must not overwrite source CSV"):
            write_purchasing_workbook(self.source, output)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(list(self.root.glob(f".{output.name}.*.xlsx")), [])

    def test_rejects_exact_source_path(self) -> None:
        self.assert_safe_rejection(self.source)

    def test_rejects_symlink_to_source(self) -> None:
        alias = self.root / "alias.xlsx"
        try:
            alias.symlink_to(self.source)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"symlinks unavailable: {exc}")
        self.assert_safe_rejection(alias)

    def test_rejects_hardlink_to_source(self) -> None:
        alias = self.root / "hardlink.xlsx"
        try:
            os.link(self.source, alias)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"hardlinks unavailable: {exc}")
        self.assert_safe_rejection(alias)

    def test_rechecks_late_symlink_before_replace(self) -> None:
        # Simulate another process creating an alias while the workbook is built.
        output = self.root / "late-alias.xlsx"
        probe = self.root / "symlink-probe"
        try:
            probe.symlink_to(self.source)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"symlinks unavailable: {exc}")
        else:
            probe.unlink()

        before = self.source.read_bytes()
        original_save = Workbook.save

        def save_and_create_alias(workbook: Workbook, filename: str | Path) -> None:
            original_save(workbook, filename)
            output.symlink_to(self.source)

        with patch("purchasing_report.Workbook.save", new=save_and_create_alias):
            with self.assertRaisesRegex(PurchasingReportError, "must not overwrite source CSV"):
                write_purchasing_workbook(self.source, output)

        self.assertEqual(self.source.read_bytes(), before)
        self.assertTrue(output.is_symlink())
        self.assertEqual(list(self.root.glob(f".{output.name}.*.xlsx")), [])

    def test_distinct_existing_workbook_can_be_replaced(self) -> None:
        output = self.root / "purchasing.xlsx"
        output.write_bytes(b"old workbook")
        write_purchasing_workbook(self.source, output)
        book = load_workbook(output, read_only=True)
        try:
            self.assertEqual(book.sheetnames, ["Summary", "Purchasing Review"])
        finally:
            book.close()
        self.assertIn(b"matched,Demo,SKU-1", self.source.read_bytes())


if __name__ == "__main__":
    unittest.main()
